#!/usr/bin/env python3
"""
Add screenshots to E-Commerce presentation with tech stack annotations
Uses only available screenshots
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os
import glob

# Paths
PRESENTATION_PATH = "/Users/rashminawale/ecommerce-app/E-Commerce-Presentation.pptx"
OUTPUT_PATH = "/Users/rashminawale/ecommerce-app/E-Commerce-Presentation-Updated.pptx"

# Find all screenshots from April 21
desktop_screenshots = glob.glob("/Users/rashminawale/Desktop/Screenshot 2026-04-21*.png")

print(f"Found {len(desktop_screenshots)} screenshots from April 21:")
for screenshot in desktop_screenshots:
    print(f"  - {os.path.basename(screenshot)}")

# Screenshot metadata - we'll match these to available files
SCREENSHOT_METADATA = {
    "6.13.55 PM": {
        "title": "Application Overview",
        "tech": "Full Stack E-Commerce Platform",
        "features": ["Angular frontend", "Spring Boot backend", "PostgreSQL database"]
    },
    "6.14.59 PM": {
        "title": "Product Listing Page",
        "tech": "Angular 17 + Angular Material",
        "features": ["Responsive product grid", "Search functionality", "Product cards", "Add to Cart buttons"]
    },
    "6.15.12 PM": {
        "title": "Shopping Cart - Water Bottle",
        "tech": "Angular Services + Spring Boot REST API",
        "features": ["Cart item management", "Quantity controls", "Order summary", "Dynamic pricing"]
    },
    "6.15.24 PM": {
        "title": "User Interface",
        "tech": "Material Design Components",
        "features": ["Clean UI design", "Responsive layout", "User-friendly navigation"]
    },
    "6.15.28 PM": {
        "title": "Empty Shopping Cart State",
        "tech": "Angular Components + State Management",
        "features": ["Empty state handling", "User-friendly messaging", "Call-to-action button"]
    },
    "6.18.00 PM": {
        "title": "Application Features",
        "tech": "Integrated Frontend & Backend",
        "features": ["Real-time updates", "Secure authentication", "RESTful APIs"]
    },
    "6.19.50 PM": {
        "title": "Product Catalog",
        "tech": "Angular + Material Design",
        "features": ["Multiple products display", "Category filtering", "Quick add to cart"]
    },
    "6.20.02 PM": {
        "title": "Product Detail Page",
        "tech": "Angular Router + Material",
        "features": ["Dynamic routing", "Product details", "Stock status", "Add to cart"]
    },
    "6.20.05 PM": {
        "title": "Shopping Cart - T-Shirt",
        "tech": "Real-time Cart Management",
        "features": ["Live cart updates", "Price calculations", "Checkout flow", "Item removal"]
    },
    "8.04.08 PM": {
        "title": "Database Configuration",
        "tech": "PostgreSQL + Spring Data JPA",
        "features": ["Database connectivity", "ORM configuration", "Connection pooling"]
    }
}

def add_tech_stack_slide(prs):
    """Add a comprehensive tech stack overview slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = "🛠️ Technology Stack Overview"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(63, 81, 181)
    title_para.alignment = PP_ALIGN.CENTER

    # Three columns
    col_width = Inches(2.8)
    col_height = Inches(5)
    start_y = Inches(1.2)

    # Frontend Column
    frontend_box = slide.shapes.add_textbox(Inches(0.5), start_y, col_width, col_height)
    frontend_frame = frontend_box.text_frame
    frontend_frame.word_wrap = True
    frontend_frame.text = """Frontend
━━━━━━━━━━━━━━
• Angular 17.3
• TypeScript 5.4
• Angular Material
• RxJS 7.8
• Angular Router
• Reactive Forms

Features:
✓ Component architecture
✓ Type-safe code
✓ Material Design UI
✓ Responsive layout
✓ Client-side routing
✓ Form validation"""

    for i, paragraph in enumerate(frontend_frame.paragraphs):
        paragraph.font.size = Pt(11)
        if i == 0:
            paragraph.font.size = Pt(18)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(33, 150, 243)

    # Backend Column
    backend_box = slide.shapes.add_textbox(Inches(3.6), start_y, col_width, col_height)
    backend_frame = backend_box.text_frame
    backend_frame.word_wrap = True
    backend_frame.text = """Backend
━━━━━━━━━━━━━━
• Spring Boot 3.2.4
• Java 17
• Spring Security
• Spring Data JPA
• JWT Auth (0.12.5)
• Maven Build

Security:
✓ JWT authentication
✓ BCrypt passwords
✓ Rate limiting
✓ XSS protection
✓ CORS config
✓ SQL injection prevention"""

    for i, paragraph in enumerate(backend_frame.paragraphs):
        paragraph.font.size = Pt(11)
        if i == 0:
            paragraph.font.size = Pt(18)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(76, 175, 80)

    # Database Column
    database_box = slide.shapes.add_textbox(Inches(6.7), start_y, col_width, col_height)
    database_frame = database_box.text_frame
    database_frame.word_wrap = True
    database_frame.text = """Database & Tools
━━━━━━━━━━━━━━
• PostgreSQL
• JPA/Hibernate ORM
• Parameterized queries
• Transaction mgmt

Architecture:
✓ RESTful API
✓ Layered design
✓ DTO pattern
✓ Service layer
✓ Repository pattern
✓ Exception handling

Endpoints:
• localhost:4200 (frontend)
• localhost:8080 (backend)"""

    for i, paragraph in enumerate(database_frame.paragraphs):
        paragraph.font.size = Pt(11)
        if i == 0:
            paragraph.font.size = Pt(18)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(255, 152, 0)

    # Footer
    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.4))
    footer_frame = footer_box.text_frame
    footer_frame.text = "🛡️ Security-First | Full-Stack | Production-Ready"
    footer_para = footer_frame.paragraphs[0]
    footer_para.font.size = Pt(14)
    footer_para.font.bold = True
    footer_para.alignment = PP_ALIGN.CENTER
    footer_para.font.color.rgb = RGBColor(96, 96, 96)

def add_screenshot_slide(prs, screenshot_path, metadata):
    """Add a slide with screenshot and tech info"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.5))
    title_frame = title_box.text_frame
    title_frame.text = metadata["title"]
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(63, 81, 181)

    # Tech label
    tech_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(9), Inches(0.25))
    tech_frame = tech_box.text_frame
    tech_frame.text = f"💻 Technology: {metadata['tech']}"
    tech_para = tech_frame.paragraphs[0]
    tech_para.font.size = Pt(14)
    tech_para.font.italic = True
    tech_para.font.color.rgb = RGBColor(76, 175, 80)

    # Add screenshot
    try:
        left = Inches(0.5)
        top = Inches(1.1)
        height = Inches(3.7)
        slide.shapes.add_picture(screenshot_path, left, top, height=height)
    except Exception as e:
        print(f"Error adding image: {e}")

    # Features
    features_box = slide.shapes.add_textbox(Inches(0.5), Inches(5), Inches(9), Inches(1.9))
    features_frame = features_box.text_frame
    features_frame.word_wrap = True
    features_text = "✨ Key Features:\n" + "\n".join([f"   • {feature}" for feature in metadata["features"]])
    features_frame.text = features_text

    for i, paragraph in enumerate(features_frame.paragraphs):
        if i == 0:
            paragraph.font.size = Pt(16)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(63, 81, 181)
        else:
            paragraph.font.size = Pt(13)
            paragraph.space_before = Pt(3)

def main():
    print("\n" + "="*60)
    print("E-Commerce Presentation Screenshot Updater")
    print("="*60 + "\n")

    print(f"Loading presentation from: {PRESENTATION_PATH}")
    prs = Presentation(PRESENTATION_PATH)
    print(f"✓ Original presentation has {len(prs.slides)} slides\n")

    # Add tech stack overview
    print("Adding Tech Stack Overview slide...")
    add_tech_stack_slide(prs)
    print("✓ Tech stack slide added\n")

    # Add screenshot slides
    added_count = 0
    for screenshot_path in desktop_screenshots:
        filename = os.path.basename(screenshot_path)

        # Extract time from filename to match metadata
        for time_key, metadata in SCREENSHOT_METADATA.items():
            if time_key in filename:
                print(f"Adding: {metadata['title']}")
                add_screenshot_slide(prs, screenshot_path, metadata)
                added_count += 1
                break

    # Save
    print(f"\n{'='*60}")
    print(f"Saving presentation to: {OUTPUT_PATH}")
    prs.save(OUTPUT_PATH)
    print(f"{'='*60}\n")
    print(f"✅ SUCCESS!")
    print(f"   • Added tech stack overview slide")
    print(f"   • Added {added_count} screenshot slides")
    print(f"   • Final presentation has {len(prs.slides)} slides")
    print(f"\n📁 File saved: E-Commerce-Presentation-Updated.pptx\n")

if __name__ == "__main__":
    main()
