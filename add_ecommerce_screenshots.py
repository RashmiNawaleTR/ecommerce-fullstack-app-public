#!/usr/bin/env python3
"""Add E-Commerce screenshots to presentation with tech stack details"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

PRESENTATION_PATH = "/Users/rashminawale/ecommerce-app/E-Commerce-Presentation.pptx"
OUTPUT_PATH = "/Users/rashminawale/ecommerce-app/E-Commerce-Presentation-Updated.pptx"

# Specific screenshots to add with metadata
SCREENSHOTS = [
    {
        "path": "/Users/rashminawale/Desktop/Screenshot 2026-04-21 at 6.19.50 PM.png",
        "title": "Product Catalog - Frontend",
        "tech": "Angular 17 + Material Design",
        "description": "Product listing with responsive grid layout showing multiple products with images, descriptions, and Add to Cart functionality."
    },
    {
        "path": "/Users/rashminawale/Desktop/Screenshot 2026-04-21 at 6.15.28 PM.png",
        "title": "Empty Cart State - UI/UX",
        "tech": "Angular Components + State Management",
        "description": "User-friendly empty cart page with clear messaging and call-to-action to browse products."
    },
    {
        "path": "/Users/rashminawale/Desktop/Screenshot 2026-04-21 at 6.20.05 PM.png",
        "title": "Shopping Cart - Full Stack Integration",
        "tech": "Angular + Spring Boot REST API",
        "description": "Active shopping cart showing item management, quantity controls, real-time price calculations, and checkout flow."
    },
    {
        "path": "/Users/rashminawale/Desktop/Screenshot 2026-04-21 at 6.15.12 PM.png",
        "title": "Cart Item Management",
        "tech": "Real-time Updates via REST API",
        "description": "Cart displaying different product with dynamic pricing and quantity management features."
    },
    {
        "path": "/Users/rashminawale/Desktop/Screenshot 2026-04-21 at 6.20.02 PM.png",
        "title": "Product Detail Page",
        "tech": "Angular Router + Dynamic Components",
        "description": "Detailed product view with category badges, stock status, pricing, and add to cart functionality."
    }
]

def add_tech_stack_overview(prs):
    """Add comprehensive tech stack slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    # Title
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    tf = title.text_frame
    tf.text = "🛠️ E-Commerce Platform - Technology Stack"
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(63, 81, 181)
    p.alignment = PP_ALIGN.CENTER

    # Frontend section
    frontend = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(2.8), Inches(5))
    tf = frontend.text_frame
    tf.word_wrap = True
    tf.text = """FRONTEND
━━━━━━━━━━━━━━

Angular 17.3
TypeScript 5.4
Angular Material
RxJS 7.8
Angular Router
Reactive Forms

FEATURES:
✓ Component-based
✓ Type-safe code
✓ Material Design
✓ Responsive UI
✓ Client routing
✓ Form validation"""

    for i, p in enumerate(tf.paragraphs):
        p.font.size = Pt(11) if i > 0 else Pt(16)
        p.font.bold = (i == 0)
        if i == 0:
            p.font.color.rgb = RGBColor(33, 150, 243)

    # Backend section
    backend = slide.shapes.add_textbox(Inches(3.6), Inches(1.2), Inches(2.8), Inches(5))
    tf = backend.text_frame
    tf.word_wrap = True
    tf.text = """BACKEND
━━━━━━━━━━━━━━

Spring Boot 3.2.4
Java 17
Spring Security
Spring Data JPA
JWT Auth
Maven

SECURITY:
✓ JWT tokens
✓ BCrypt hashing
✓ Rate limiting
✓ XSS protection
✓ CORS config
✓ SQL injection protection"""

    for i, p in enumerate(tf.paragraphs):
        p.font.size = Pt(11) if i > 0 else Pt(16)
        p.font.bold = (i == 0)
        if i == 0:
            p.font.color.rgb = RGBColor(76, 175, 80)

    # Database & Architecture section
    database = slide.shapes.add_textbox(Inches(6.7), Inches(1.2), Inches(2.8), Inches(5))
    tf = database.text_frame
    tf.word_wrap = True
    tf.text = """DATABASE
━━━━━━━━━━━━━━

PostgreSQL
JPA/Hibernate ORM
Parameterized queries
Transaction mgmt

ARCHITECTURE:
✓ RESTful API
✓ Layered design
✓ DTO pattern
✓ Service layer
✓ Repository pattern

DEPLOYMENT:
• localhost:4200 (UI)
• localhost:8080 (API)"""

    for i, p in enumerate(tf.paragraphs):
        p.font.size = Pt(11) if i > 0 else Pt(16)
        p.font.bold = (i == 0)
        if i == 0:
            p.font.color.rgb = RGBColor(255, 152, 0)

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.4))
    tf = footer.text_frame
    tf.text = "🛡️ Security-First Development | Full-Stack Integration | Production-Ready"
    p = tf.paragraphs[0]
    p.font.size = Pt(13)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(80, 80, 80)

def add_screenshot_slide(prs, screenshot):
    """Add slide with screenshot and details"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    # Title
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.45))
    tf = title.text_frame
    tf.text = screenshot["title"]
    p = tf.paragraphs[0]
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(63, 81, 181)

    # Tech tag
    tech = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(9), Inches(0.25))
    tf = tech.text_frame
    tf.text = f"💻 {screenshot['tech']}"
    p = tf.paragraphs[0]
    p.font.size = Pt(13)
    p.font.italic = True
    p.font.color.rgb = RGBColor(76, 175, 80)

    # Screenshot image
    if os.path.exists(screenshot["path"]):
        slide.shapes.add_picture(
            screenshot["path"],
            Inches(0.5),
            Inches(1.1),
            height=Inches(3.6)
        )

    # Description
    desc = slide.shapes.add_textbox(Inches(0.5), Inches(4.9), Inches(9), Inches(1.8))
    tf = desc.text_frame
    tf.word_wrap = True
    tf.text = f"IMPLEMENTATION DETAILS\n\n{screenshot['description']}"

    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(63, 81, 181)

    if len(tf.paragraphs) > 1:
        tf.paragraphs[2].font.size = Pt(12)
        tf.paragraphs[2].space_before = Pt(6)

def main():
    print("\n" + "="*70)
    print("E-COMMERCE PRESENTATION - ADDING SCREENSHOTS WITH TECH STACK")
    print("="*70 + "\n")

    print(f"📂 Loading: {os.path.basename(PRESENTATION_PATH)}")
    prs = Presentation(PRESENTATION_PATH)
    print(f"   Current slides: {len(prs.slides)}\n")

    # Add tech stack overview
    print("📊 Adding Tech Stack Overview...")
    add_tech_stack_overview(prs)
    print("   ✓ Tech stack slide added\n")

    # Add screenshot slides
    print("📸 Adding Application Screenshots:\n")
    added = 0
    for screenshot in SCREENSHOTS:
        if os.path.exists(screenshot["path"]):
            print(f"   ✓ {screenshot['title']}")
            add_screenshot_slide(prs, screenshot)
            added += 1
        else:
            print(f"   ✗ Not found: {screenshot['title']}")

    # Save
    print(f"\n{'='*70}")
    prs.save(OUTPUT_PATH)
    print(f"💾 SAVED: {os.path.basename(OUTPUT_PATH)}")
    print(f"{'='*70}\n")

    print(f"✅ SUMMARY:")
    print(f"   • Tech Stack Overview: 1 slide")
    print(f"   • Screenshots Added: {added} slides")
    print(f"   • Total Slides: {len(prs.slides)}")
    print(f"\n📁 Location: {OUTPUT_PATH}\n")

if __name__ == "__main__":
    main()
